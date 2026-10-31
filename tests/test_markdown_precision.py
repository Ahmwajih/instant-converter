import json
from pathlib import Path
from docx import Document
import openpyxl
import pymupdf
from pptx import Presentation
from converter.engine import convert_file


def test_pdf_to_markdown_preserves_headings_bold_and_tables(tmp_path: Path) -> None:
    pdf_path = tmp_path / "architecture_spec.pdf"
    doc = pymupdf.open()
    page = doc.new_page()

    page.insert_text((50, 60), "System Architecture Guide", fontsize=22, fontname="helv")
    page.insert_text((50, 95), "Database Layer Overview", fontsize=15, fontname="helv")
    page.insert_text((50, 125), "The primary storage engine uses PostgreSQL for ACID transactions.", fontsize=11, fontname="helv")
    page.insert_text((50, 162), "• Automated daily snapshots", fontsize=11, fontname="helv")

    x0, y0, col_w, row_h = 50, 200, 140, 26
    for r in range(3):
        page.draw_line((x0, y0 + r * row_h), (x0 + 2 * col_w, y0 + r * row_h))
    for c in range(3):
        page.draw_line((x0 + c * col_w, y0), (x0 + c * col_w, y0 + 2 * row_h))

    page.insert_text((x0 + 8, y0 + 18), "Service", fontsize=11, fontname="helv")
    page.insert_text((x0 + col_w + 8, y0 + 18), "Port", fontsize=11, fontname="helv")
    page.insert_text((x0 + 8, y0 + row_h + 18), "API Gateway", fontsize=11, fontname="helv")
    page.insert_text((x0 + col_w + 8, y0 + row_h + 18), "8443", fontsize=11, fontname="helv")

    doc.save(str(pdf_path))
    doc.close()

    md_out = convert_file(pdf_path, "MD")
    md_text = md_out.read_text(encoding="utf-8")

    assert "# System Architecture Guide" in md_text
    assert "## Database Layer Overview" in md_text
    assert "- Automated daily snapshots" in md_text
    assert "| Service | Port |" in md_text
    assert "| API Gateway | 8443 |" in md_text


def test_docx_to_markdown_preserves_headings_runs_lists_and_tables(tmp_path: Path) -> None:
    docx_path = tmp_path / "runbook.docx"
    doc = Document()
    doc.add_heading("Deployment Runbook", level=1)
    doc.add_heading("Prerequisites", level=2)

    p = doc.add_paragraph("Always verify ")
    r_bold = p.add_run("zero-downtime")
    r_bold.bold = True
    p.add_run(" migrations and ")
    r_italic = p.add_run("backup integrity")
    r_italic.italic = True
    p.add_run(".")

    doc.add_paragraph("Verify SSH keys", style="List Bullet")

    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "Environment"
    tbl.cell(0, 1).text = "Region"
    tbl.cell(1, 0).text = "Production"
    tbl.cell(1, 1).text = "eu-central-1"

    doc.save(str(docx_path))

    md_out = convert_file(docx_path, "MD")
    md_text = md_out.read_text(encoding="utf-8")

    assert "# Deployment Runbook" in md_text
    assert "## Prerequisites" in md_text
    assert "**zero-downtime**" in md_text
    assert "*backup integrity*" in md_text
    assert "- Verify SSH keys" in md_text
    assert "| Environment | Region |" in md_text
    assert "| Production | eu-central-1 |" in md_text


def test_html_pptx_and_xlsx_to_markdown(tmp_path: Path) -> None:
    html_path = tmp_path / "guide.html"
    html_path.write_text(
        "<h1>API Docs</h1><p>Use <code>GET /v1/health</code> with <strong>Bearer</strong> auth.</p>"
        "<table><tr><th>Code</th><th>Status</th></tr><tr><td>200</td><td>OK</td></tr></table>",
        encoding="utf-8",
    )
    md_from_html = convert_file(html_path, "MD").read_text(encoding="utf-8")
    assert "# API Docs" in md_from_html
    assert "`GET /v1/health`" in md_from_html
    assert "**Bearer**" in md_from_html
    assert "| Code | Status |" in md_from_html

    pptx_path = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Q4 Roadmap"
    slide.placeholders[1].text = "Launch v1.0.2"
    prs.save(str(pptx_path))
    md_from_pptx = convert_file(pptx_path, "MD").read_text(encoding="utf-8")
    assert "## Slide 1: Q4 Roadmap" in md_from_pptx
    assert "- Launch v1.0.2" in md_from_pptx

    xlsx_path = tmp_path / "metrics.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Latency"
    ws.append(["Endpoint", "p99_ms"])
    ws.append(["/convert", 12])
    wb.save(str(xlsx_path))
    md_from_xlsx = convert_file(xlsx_path, "MD").read_text(encoding="utf-8")
    assert "## Latency" in md_from_xlsx
    assert "| Endpoint | p99_ms |" in md_from_xlsx
    assert "| /convert | 12 |" in md_from_xlsx


def test_trivy_json_and_nested_json_to_readable_markdown(tmp_path: Path) -> None:
    trivy_file = tmp_path / "trivy_harbor_scan.json"
    cve_url = "https://" + "avd.aquasec.com/nvd/cve-2026-1234"
    trivy_payload = {
        "SchemaVersion": 2,
        "ArtifactName": "harbor.local/prod/backend:1.4.2",
        "ArtifactType": "container_image",
        "Results": [
            {
                "Target": "harbor.local/prod/backend:1.4.2 (debian 12.5)",
                "Class": "os-pkgs",
                "Type": "debian",
                "Vulnerabilities": [
                    {
                        "VulnerabilityID": "CVE-2026-1234",
                        "PkgName": "openssl",
                        "InstalledVersion": "3.0.11-1",
                        "FixedVersion": "3.0.14-1",
                        "Severity": "CRITICAL",
                        "Title": "Remote buffer overflow in X.509 certificate verification",
                        "PrimaryURL": cve_url,
                    }
                ],
            }
        ],
    }
    trivy_file.write_text(json.dumps(trivy_payload, indent=2), encoding="utf-8")

    md_out = convert_file(trivy_file, "MD")
    md_text = md_out.read_text(encoding="utf-8")

    assert "# Security Scan Report: `harbor.local/prod/backend:1.4.2`" in md_text
    assert "## Executive Severity Summary" in md_text
    assert "| CRITICAL | 1 |" in md_text
    assert f"[CVE-2026-1234]({cve_url})" in md_text
    assert "`openssl`" in md_text
    assert "`3.0.14-1`" in md_text
