from collections import Counter
from html.parser import HTMLParser
import csv
import re
import shutil
import subprocess
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from typing import Any
import openpyxl
import pymupdf


def _clean_cell(val: Any) -> str:
    if val is None:
        return ""
    text = str(val).replace("\r\n", " ").replace("\n", " ").strip()
    return text.replace("|", "\\|")


def rows_to_markdown_table(rows: list[list[Any]]) -> str:
    """Formats a 2D list of rows into a clean GitHub-Flavored Markdown table."""
    cleaned = [[_clean_cell(c) for c in row] for row in rows if any(_clean_cell(c) for c in row)]
    if not cleaned:
        return ""
    col_count = max(len(r) for r in cleaned)
    normalized = [r + [""] * (col_count - len(r)) for r in cleaned]
    keep_cols = [c for c in range(col_count) if any(row[c] != "" for row in normalized)]
    if keep_cols:
        normalized = [[row[c] for c in keep_cols] for row in normalized]
        col_count = len(keep_cols)
    header = normalized[0]
    sep = ["---"] * col_count
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(sep) + " |",
    ]
    for row in normalized[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


_BULLET_PREFIX_RE = re.compile(r"^([\u00b7\u2022\u2023\u25E6\u2043\u2219\u25aa\u25cf\uf0b7\*\-\+])\s+(.*)$")
_NUMBER_PREFIX_RE = re.compile(r"^(\d+[\.\)])\s+(.*)$")


def _rect_intersects(r1: tuple[float, float, float, float], r2: tuple[float, float, float, float]) -> bool:
    ix0 = max(r1[0], r2[0])
    iy0 = max(r1[1], r2[1])
    ix1 = min(r1[2], r2[2])
    iy1 = min(r1[3], r2[3])
    return ix1 > ix0 and iy1 > iy0


def _format_pdf_span(text: str, flags: int, font_name: str, is_heading: bool = False) -> str:
    if not text.strip():
        return text
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()) :]
    core = text.strip()

    fn_lower = font_name.lower()
    is_mono = bool(flags & 8) or any(k in fn_lower for k in ("mono", "courier", "consolas", "menlo", "code"))
    is_bold = bool(flags & 16) or any(k in fn_lower for k in ("bold", "black", "heavy", "semibold", "demi"))
    is_italic = bool(flags & 2) or any(k in fn_lower for k in ("italic", "oblique", "slant"))

    if is_mono and not is_heading:
        return f"{lead}`{core}`{trail}"
    if is_heading:
        return f"{lead}{core}{trail}"
    if is_bold and is_italic:
        return f"{lead}***{core}***{trail}"
    if is_bold:
        return f"{lead}**{core}**{trail}"
    if is_italic:
        return f"{lead}*{core}*{trail}"
    return text


def pdf_to_markdown(source_path: Path, output_path: Path) -> Path:
    doc = pymupdf.open(str(source_path))
    if len(doc) == 0:
        doc.close()
        output_path.write_text("", encoding="utf-8")
        return output_path

    size_counter: Counter[float] = Counter()
    for page in doc:
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    txt = span.get("text", "").strip()
                    if txt:
                        sz = round(float(span.get("size", 11.0)), 1)
                        size_counter[sz] += len(txt)

    body_size = size_counter.most_common(1)[0][0] if size_counter else 11.0
    md_pages: list[str] = []

    for page in doc:
        elements: list[tuple[float, float, str]] = []
        table_bboxes: list[tuple[float, float, float, float]] = []

        try:
            found_tables = page.find_tables()
            for tab in getattr(found_tables, "tables", found_tables):
                bbox = tuple(float(v) for v in tab.bbox)
                table_bboxes.append(bbox)
                md_tab = ""
                extracted = tab.extract()
                if extracted:
                    md_tab = rows_to_markdown_table(extracted)
                elif hasattr(tab, "to_markdown"):
                    md_tab = (tab.to_markdown() or "").strip()
                if md_tab:
                    elements.append((bbox[1], bbox[0], f"\n{md_tab}\n"))
        except Exception:
            pass

        links: list[tuple[tuple[float, float, float, float], str]] = []
        try:
            for lnk in page.get_links():
                uri = lnk.get("uri")
                from_rect = lnk.get("from")
                if uri and from_rect:
                    links.append(((float(from_rect.x0), float(from_rect.y0), float(from_rect.x1), float(from_rect.y1)), uri))
        except Exception:
            pass

        text_dict = page.get_text("dict")
        page_has_text = False

        for block in text_dict.get("blocks", []):
            if block.get("type") != 0:
                continue
            b_bbox = tuple(float(v) for v in block.get("bbox", (0, 0, 0, 0)))
            if any(_rect_intersects(b_bbox, tb) for tb in table_bboxes):
                continue

            block_lines: list[str] = []
            for line in block.get("lines", []):
                l_bbox = tuple(float(v) for v in line.get("bbox", (0, 0, 0, 0)))
                if any(_rect_intersects(l_bbox, tb) for tb in table_bboxes):
                    continue

                spans = line.get("spans", [])
                if not spans:
                    continue

                max_sz = max(float(s.get("size", body_size)) for s in spans)
                ratio = max_sz / max(1.0, body_size)
                heading_prefix = ""
                if ratio >= 1.55 or max_sz >= 20.0:
                    heading_prefix = "# "
                elif ratio >= 1.30 or max_sz >= 16.0:
                    heading_prefix = "## "
                elif ratio >= 1.15 or max_sz >= 13.5:
                    heading_prefix = "### "

                is_heading = bool(heading_prefix)
                rendered_spans: list[str] = []
                for s in spans:
                    raw_t = s.get("text", "")
                    if not raw_t:
                        continue
                    page_has_text = True
                    s_bbox = tuple(float(v) for v in s.get("bbox", l_bbox))
                    formatted = _format_pdf_span(
                        raw_t,
                        int(s.get("flags", 0)),
                        str(s.get("font", "")),
                        is_heading=is_heading,
                    )
                    for l_rect, uri in links:
                        if _rect_intersects(s_bbox, l_rect) and formatted.strip():
                            formatted = f"[{formatted.strip()}]({uri})"
                            break
                    rendered_spans.append(formatted)

                line_str = "".join(rendered_spans).strip()
                if not line_str:
                    continue

                line_str = line_str.replace("****", "").replace("** **", " ")

                if heading_prefix:
                    block_lines.append(f"\n{heading_prefix}{line_str}\n")
                else:
                    m_bul = _BULLET_PREFIX_RE.match(line_str)
                    m_num = _NUMBER_PREFIX_RE.match(line_str)
                    if m_bul:
                        block_lines.append(f"- {m_bul.group(2)}")
                    elif m_num:
                        block_lines.append(f"{m_num.group(1)} {m_num.group(2)}")
                    else:
                        block_lines.append(line_str)

            if block_lines:
                joined_block = "\n".join(block_lines).strip()
                elements.append((b_bbox[1], b_bbox[0], joined_block))

        if not page_has_text and not elements:
            try:
                from converter.images import _extract_text_from_image
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_img:
                    tmp_png = Path(tmp_img.name)
                page.get_pixmap(dpi=200, alpha=False).save(str(tmp_png))
                ocr_txt = _extract_text_from_image(tmp_png).strip()
                tmp_png.unlink(missing_ok=True)
                if ocr_txt:
                    elements.append((0.0, 0.0, ocr_txt))
            except Exception:
                pass

        elements.sort(key=lambda item: (round(item[0] / 8.0), item[1]))
        page_md = "\n\n".join(el[2] for el in elements if el[2].strip())
        if page_md.strip():
            md_pages.append(page_md.strip())

    doc.close()
    final_md = "\n\n---\n\n".join(md_pages)
    final_md = re.sub(r"\n{3,}", "\n\n", final_md).strip() + "\n"
    output_path.write_text(final_md, encoding="utf-8")
    return output_path


_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _format_docx_run(r_elem: ET.Element, is_heading: bool = False) -> str:
    texts = [t.text for t in r_elem.iter(f"{_W_NS}t") if t.text]
    raw = "".join(texts)
    if not raw or not raw.strip():
        return raw

    rPr = r_elem.find(f"{_W_NS}rPr")
    is_bold = False
    is_italic = False
    is_strike = False
    is_mono = False

    if rPr is not None:
        b_el = rPr.find(f"{_W_NS}b")
        i_el = rPr.find(f"{_W_NS}i")
        s_el = rPr.find(f"{_W_NS}strike")
        f_el = rPr.find(f"{_W_NS}rFonts")
        is_bold = b_el is not None and b_el.get(f"{_W_NS}val", "true").lower() not in ("false", "0")
        is_italic = i_el is not None and i_el.get(f"{_W_NS}val", "true").lower() not in ("false", "0")
        is_strike = s_el is not None and s_el.get(f"{_W_NS}val", "true").lower() not in ("false", "0")
        if f_el is not None:
            font_ascii = (f_el.get(f"{_W_NS}ascii") or "").lower()
            is_mono = any(k in font_ascii for k in ("courier", "consolas", "mono", "menlo"))

    lead = raw[: len(raw) - len(raw.lstrip())]
    trail = raw[len(raw.rstrip()) :]
    core = raw.strip()

    if is_mono and not is_heading:
        core = f"`{core}`"
    elif not is_heading:
        if is_bold and is_italic:
            core = f"***{core}***"
        elif is_bold:
            core = f"**{core}**"
        elif is_italic:
            core = f"*{core}*"
    if is_strike:
        core = f"~~{core}~~"
    return f"{lead}{core}{trail}"


def _docx_paragraph_to_md(p_elem: ET.Element, rels: dict[str, str], styles_map: dict[str, str]) -> str:
    style_id = ""
    is_list = False
    list_level = 0

    pPr = p_elem.find(f"{_W_NS}pPr")
    if pPr is not None:
        pStyle = pPr.find(f"{_W_NS}pStyle")
        if pStyle is not None:
            style_id = pStyle.get(f"{_W_NS}val", "")
        numPr = pPr.find(f"{_W_NS}numPr")
        if numPr is not None:
            is_list = True
            ilvl = numPr.find(f"{_W_NS}ilvl")
            if ilvl is not None:
                try:
                    list_level = int(ilvl.get(f"{_W_NS}val", "0"))
                except ValueError:
                    list_level = 0

    style_name = styles_map.get(style_id, style_id).lower()
    heading_prefix = ""
    if "title" in style_name or style_name in ("heading 1", "heading1", "h1"):
        heading_prefix = "# "
    elif style_name in ("heading 2", "heading2", "h2", "subtitle"):
        heading_prefix = "## "
    elif style_name in ("heading 3", "heading3", "h3"):
        heading_prefix = "### "
    elif style_name in ("heading 4", "heading4", "h4"):
        heading_prefix = "#### "
    elif style_name in ("heading 5", "heading5", "h5"):
        heading_prefix = "##### "
    elif style_name in ("heading 6", "heading6", "h6"):
        heading_prefix = "###### "

    is_heading = bool(heading_prefix)
    parts: list[str] = []

    for child in p_elem:
        if child.tag == f"{_W_NS}r":
            parts.append(_format_docx_run(child, is_heading=is_heading))
        elif child.tag == f"{_W_NS}hyperlink":
            r_id = child.get(f"{_R_NS}id", "")
            link_runs = [_format_docx_run(r, is_heading=True) for r in child.findall(f"{_W_NS}r")]
            label = "".join(link_runs).strip()
            url = rels.get(r_id, "")
            if label and url:
                parts.append(f"[{label}]({url})")
            elif label:
                parts.append(label)

    line = "".join(parts).strip()
    if not line:
        return ""
    line = line.replace("****", "").replace("** **", " ")

    if heading_prefix:
        return f"{heading_prefix}{line}"
    if "quote" in style_name:
        return f"> {line}"
    if is_list or "list" in style_name or "bullet" in style_name:
        indent = "  " * max(0, list_level)
        if "number" in style_name:
            return f"{indent}1. {line}"
        return f"{indent}- {line}"
    return line


def docx_to_markdown(source_path: Path, output_path: Path) -> Path:
    from docx import Document

    doc = Document(str(source_path))
    rels: dict[str, str] = {}
    for rel_id, rel in doc.part.rels.items():
        if "hyperlink" in rel.reltype:
            rels[rel_id] = rel._target

    styles_map: dict[str, str] = {}
    for s in doc.styles:
        if getattr(s, "style_id", None) and getattr(s, "name", None):
            styles_map[s.style_id] = s.name

    blocks: list[str] = []
    body = doc.element.body
    for child in body:
        if child.tag == f"{_W_NS}p":
            md_p = _docx_paragraph_to_md(child, rels, styles_map)
            if md_p:
                blocks.append(md_p)
        elif child.tag == f"{_W_NS}tbl":
            rows_data: list[list[str]] = []
            for tr in child.findall(f"{_W_NS}tr"):
                row_cells: list[str] = []
                for tc in tr.findall(f"{_W_NS}tc"):
                    cell_paras = [
                        _docx_paragraph_to_md(p, rels, styles_map)
                        for p in tc.findall(f"{_W_NS}p")
                    ]
                    row_cells.append(" ".join(cp for cp in cell_paras if cp).strip())
                if any(row_cells):
                    rows_data.append(row_cells)
            if rows_data:
                blocks.append(rows_to_markdown_table(rows_data))

    final_md = "\n\n".join(blocks).strip() + "\n"
    output_path.write_text(final_md, encoding="utf-8")
    return output_path


class _HTMLToMarkdownParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.out: list[str] = []
        self.list_stack: list[str] = []
        self.in_pre = False
        self.in_code = False
        self.in_table = False
        self.table_rows: list[list[str]] = []
        self.current_row: list[str] = []
        self.current_cell: list[str] = []
        self.link_href: str | None = None
        self.link_text: list[str] = []

    def _emit(self, text: str) -> None:
        if self.link_href is not None:
            self.link_text.append(text)
        elif self.in_table and self.current_cell is not None:
            self.current_cell.append(text)
        else:
            self.out.append(text)

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        t = tag.lower()
        attr_dict = dict(attrs)
        if t in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(t[1])
            self._emit(f"\n\n{'#' * level} ")
        elif t == "p":
            self._emit("\n\n")
        elif t == "br":
            self._emit("\n")
        elif t == "hr":
            self._emit("\n\n---\n\n")
        elif t in ("strong", "b"):
            self._emit("**")
        elif t in ("em", "i"):
            self._emit("*")
        elif t in ("del", "s", "strike"):
            self._emit("~~")
        elif t == "blockquote":
            self._emit("\n\n> ")
        elif t == "pre":
            self.in_pre = True
            self._emit("\n\n```\n")
        elif t == "code":
            self.in_code = True
            if not self.in_pre:
                self._emit("`")
        elif t == "a":
            self.link_href = attr_dict.get("href") or ""
            self.link_text = []
        elif t in ("ul", "ol"):
            self.list_stack.append(t)
            self._emit("\n")
        elif t == "li":
            indent = "  " * max(0, len(self.list_stack) - 1)
            prefix = "1. " if (self.list_stack and self.list_stack[-1] == "ol") else "- "
            self._emit(f"\n{indent}{prefix}")
        elif t == "table":
            self.in_table = True
            self.table_rows = []
        elif t == "tr" and self.in_table:
            self.current_row = []
        elif t in ("th", "td") and self.in_table:
            self.current_cell = []

    def handle_endtag(self, tag: str) -> None:
        t = tag.lower()
        if t in ("h1", "h2", "h3", "h4", "h5", "h6", "p", "blockquote"):
            self._emit("\n\n")
        elif t in ("strong", "b"):
            self._emit("**")
        elif t in ("em", "i"):
            self._emit("*")
        elif t in ("del", "s", "strike"):
            self._emit("~~")
        elif t == "code":
            if not self.in_pre:
                self._emit("`")
            self.in_code = False
        elif t == "pre":
            self._emit("\n```\n\n")
            self.in_pre = False
        elif t == "a":
            label = "".join(self.link_text).strip()
            href = self.link_href or ""
            self.link_href = None
            self.link_text = []
            if label and href:
                self._emit(f"[{label}]({href})")
            elif label:
                self._emit(label)
        elif t in ("ul", "ol"):
            if self.list_stack:
                self.list_stack.pop()
            self._emit("\n")
        elif t in ("th", "td") and self.in_table:
            self.current_row.append("".join(self.current_cell).strip())
        elif t == "tr" and self.in_table:
            if any(self.current_row):
                self.table_rows.append(self.current_row)
        elif t == "table" and self.in_table:
            self.in_table = False
            if self.table_rows:
                self.out.append("\n\n" + rows_to_markdown_table(self.table_rows) + "\n\n")

    def handle_data(self, data: str) -> None:
        if not data:
            return
        if self.in_pre:
            self._emit(data)
        else:
            collapsed = re.sub(r"[ \t\r\n]+", " ", data)
            self._emit(collapsed)


def html_to_markdown(html_text: str) -> str:
    cleaned = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", html_text, flags=re.DOTALL | re.IGNORECASE)
    parser = _HTMLToMarkdownParser()
    parser.feed(cleaned)
    raw = "".join(parser.out)
    lines = [ln.rstrip() for ln in raw.splitlines()]
    compact = "\n".join(lines)
    compact = re.sub(r"\n{3,}", "\n\n", compact).strip()
    return compact + "\n"


def pptx_to_markdown(source_path: Path, output_path: Path) -> Path:
    from pptx import Presentation

    prs = Presentation(str(source_path))
    slides_md: list[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        title_txt = ""
        if slide.shapes.title and slide.shapes.title.text.strip():
            title_txt = slide.shapes.title.text.strip()
        header = f"## Slide {idx}: {title_txt}" if title_txt else f"## Slide {idx}"
        parts: list[str] = [header]

        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_table", False):
                t_rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                md_tbl = rows_to_markdown_table(t_rows)
                if md_tbl:
                    parts.append(md_tbl)
            elif getattr(shape, "has_text_frame", False):
                bullets: list[str] = []
                for para in shape.text_frame.paragraphs:
                    txt = para.text.strip()
                    if not txt:
                        continue
                    indent = "  " * int(getattr(para, "level", 0) or 0)
                    bullets.append(f"{indent}- {txt}")
                if bullets:
                    parts.append("\n".join(bullets))

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"> **Speaker Notes:** {notes}")

        slides_md.append("\n\n".join(parts))

    output_path.write_text("\n\n---\n\n".join(slides_md).strip() + "\n", encoding="utf-8")
    return output_path


def _format_trivy_json_to_markdown(data: dict[str, Any]) -> str:
    """Converts a Trivy vulnerability/misconfig JSON report into a clean executive Markdown report."""
    artifact = data.get("ArtifactName", "Scanned Target")
    art_type = data.get("ArtifactType", "unknown")
    created = data.get("CreatedAt", "")

    lines: list[str] = [
        f"# Security Scan Report: `{artifact}`",
        "",
        f"- **Artifact Type:** `{art_type}`",
    ]
    if created:
        lines.append(f"- **Scan Timestamp:** `{created}`")

    results = data.get("Results", [])
    sev_counts: Counter[str] = Counter()
    for res in results:
        for v in res.get("Vulnerabilities") or []:
            sev_counts[str(v.get("Severity", "UNKNOWN")).upper()] += 1
        for m in res.get("Misconfigurations") or []:
            sev_counts[str(m.get("Severity", "UNKNOWN")).upper()] += 1
        for s in res.get("Secrets") or []:
            sev_counts[str(s.get("Severity", "UNKNOWN")).upper()] += 1

    lines.extend(["", "## Executive Severity Summary", ""])
    summary_rows = [
        ["Severity", "Count"],
        ["CRITICAL", str(sev_counts.get("CRITICAL", 0))],
        ["HIGH", str(sev_counts.get("HIGH", 0))],
        ["MEDIUM", str(sev_counts.get("MEDIUM", 0))],
        ["LOW", str(sev_counts.get("LOW", 0))],
        ["UNKNOWN", str(sev_counts.get("UNKNOWN", 0))],
    ]
    lines.append(rows_to_markdown_table(summary_rows))

    sev_order = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3, "UNKNOWN": 4}

    for res in results:
        target = res.get("Target", "Target")
        res_class = res.get("Class", "")
        res_type = res.get("Type", "")
        vulns = res.get("Vulnerabilities") or []
        misconfigs = res.get("Misconfigurations") or []
        secrets = res.get("Secrets") or []

        if not vulns and not misconfigs and not secrets:
            continue

        subtitle = f" ({res_type})" if res_type else (f" ({res_class})" if res_class else "")
        lines.extend(["", f"## Target: `{target}`{subtitle}", ""])

        if vulns:
            sorted_vulns = sorted(
                vulns,
                key=lambda x: sev_order.get(str(x.get("Severity", "UNKNOWN")).upper(), 9),
            )
            v_rows: list[list[str]] = [
                ["CVE / ID", "Severity", "Package", "Installed", "Fixed In", "Title"]
            ]
            for v in sorted_vulns:
                vid = str(v.get("VulnerabilityID", ""))
                url = str(v.get("PrimaryURL", ""))
                id_cell = f"[{vid}]({url})" if url else vid
                sev = str(v.get("Severity", "UNKNOWN")).upper()
                pkg = str(v.get("PkgName", ""))
                inst = str(v.get("InstalledVersion", ""))
                fix = str(v.get("FixedVersion", "") or "—")
                title = str(v.get("Title") or v.get("Description") or "").strip()
                if len(title) > 90:
                    title = title[:87] + "..."
                sev_label = f"**{sev}**" if sev in ("CRITICAL", "HIGH") else sev
                v_rows.append([id_cell, sev_label, f"`{pkg}`", f"`{inst}`", f"`{fix}`", title])
            lines.append(rows_to_markdown_table(v_rows))

        if misconfigs:
            lines.extend(["", "### Misconfigurations", ""])
            m_rows: list[list[str]] = [["ID", "Severity", "Title", "Resolution"]]
            for m in misconfigs:
                mid = str(m.get("ID", ""))
                sev = str(m.get("Severity", "UNKNOWN")).upper()
                title = str(m.get("Title", ""))[:80]
                res_txt = str(m.get("Resolution", ""))[:90]
                m_rows.append([f"`{mid}`", sev, title, res_txt])
            lines.append(rows_to_markdown_table(m_rows))

        if secrets:
            lines.extend(["", "### Exposed Secrets", ""])
            s_rows: list[list[str]] = [["Rule ID", "Category", "Severity", "Title", "Lines"]]
            for s in secrets:
                rid = str(s.get("RuleID", ""))
                cat = str(s.get("Category", ""))
                sev = str(s.get("Severity", "UNKNOWN")).upper()
                title = str(s.get("Title", ""))[:80]
                span = f"{s.get('StartLine', '?')}-{s.get('EndLine', '?')}"
                s_rows.append([f"`{rid}`", cat, sev, title, span])
            lines.append(rows_to_markdown_table(s_rows))

    return "\n".join(lines).strip() + "\n"


def _compact_json_cell(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, (str, int, float, bool)):
        s = _clean_cell(val)
        return (s[:95] + "...") if len(s) > 98 else s
    if isinstance(val, list):
        if not val:
            return ""
        if all(isinstance(x, (str, int, float, bool)) for x in val):
            joined = ", ".join(_clean_cell(x) for x in val[:6])
            if len(val) > 6:
                joined += f" (+{len(val) - 6} more)"
            return joined
        return f"{len(val)} items"
    if isinstance(val, dict):
        pairs = [f"{k}: {_clean_cell(v)}" for k, v in list(val.items())[:4] if not isinstance(v, (dict, list))]
        return "; ".join(pairs) if pairs else f"{len(val)} fields"
    return _clean_cell(val)


def _nested_data_to_markdown(data: Any, title: str = "Document", level: int = 1) -> str:
    if isinstance(data, dict) and ("Results" in data and ("ArtifactName" in data or "SchemaVersion" in data)):
        return _format_trivy_json_to_markdown(data)

    out: list[str] = []
    if level == 1:
        out.append(f"# {title}\n")

    if isinstance(data, list):
        if not data:
            return ""
        if all(isinstance(item, dict) for item in data):
            headers: list[str] = []
            for item in data:
                for k in item.keys():
                    if k not in headers:
                        headers.append(k)
            headers = headers[:10]
            rows = [headers] + [[_compact_json_cell(item.get(h)) for h in headers] for item in data]
            out.append(rows_to_markdown_table(rows))
        elif all(isinstance(item, (str, int, float, bool)) for item in data):
            for item in data:
                out.append(f"- {_clean_cell(item)}")
        else:
            for idx, item in enumerate(data, start=1):
                out.append(_nested_data_to_markdown(item, title=f"Item {idx}", level=min(6, level + 1)))
        return "\n\n".join(p for p in out if p).strip()

    if isinstance(data, dict):
        scalars: list[str] = []
        complex_items: list[tuple[str, Any]] = []
        for k, v in data.items():
            if v is None or v == "" or v == [] or v == {}:
                continue
            if isinstance(v, (str, int, float, bool)):
                scalars.append(f"- **{k}:** {_clean_cell(v)}")
            elif isinstance(v, list) and all(isinstance(x, (str, int, float, bool)) for x in v):
                joined = ", ".join(f"`{_clean_cell(x)}`" for x in v[:12])
                if len(v) > 12:
                    joined += f" (+{len(v) - 12} more)"
                scalars.append(f"- **{k}:** {joined}")
            else:
                complex_items.append((str(k), v))

        if scalars:
            out.append("\n".join(scalars))

        for k, v in complex_items:
            hdr_lvl = "#" * min(6, level + 1)
            out.append(f"{hdr_lvl} {k}")
            out.append(_nested_data_to_markdown(v, title=k, level=min(6, level + 1)))

        return "\n\n".join(p for p in out if p).strip()

    return _clean_cell(data)


def spreadsheet_to_markdown(source_path: Path, output_path: Path) -> Path:
    ext = source_path.suffix.lower().lstrip(".")
    if ext in {"csv", "tsv"}:
        delim = "\t" if ext == "tsv" else ","
        with source_path.open("r", encoding="utf-8", errors="replace", newline="") as f:
            rows = list(csv.reader(f, delimiter=delim))
        output_path.write_text(rows_to_markdown_table(rows) + "\n", encoding="utf-8")
        return output_path

    if ext == "xlsx":
        wb = openpyxl.load_workbook(str(source_path), data_only=True)
        sections: list[str] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            md_tbl = rows_to_markdown_table(rows)
            if md_tbl:
                sections.append(f"## {sheet_name}\n\n{md_tbl}")
        wb.close()
        output_path.write_text("\n\n".join(sections).strip() + "\n", encoding="utf-8")
        return output_path

    if ext in {"json", "yaml", "yml", "toml", "xml"}:
        import json
        import tomllib
        import yaml
        raw_text = source_path.read_text(encoding="utf-8", errors="replace")
        if ext == "json":
            data = json.loads(raw_text)
        elif ext in {"yaml", "yml"}:
            data = yaml.safe_load(raw_text)
        elif ext == "toml":
            data = tomllib.loads(raw_text)
        else:
            root = ET.fromstring(raw_text)
            data = {root.tag: [{child.tag: (child.text or "").strip()} for child in root]}
        md_report = _nested_data_to_markdown(data, title=source_path.stem)
        output_path.write_text(md_report.strip() + "\n", encoding="utf-8")
        return output_path

    output_path.write_text(source_path.read_text(encoding="utf-8", errors="replace"), encoding="utf-8")
    return output_path


def odt_or_office_to_markdown(source_path: Path, output_path: Path) -> Path:
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if soffice:
        with tempfile.TemporaryDirectory() as tmpdir:
            profile_uri = Path(tmpdir).as_uri() + "/lo_profile"
            cmd = [
                soffice,
                f"-env:UserInstallation={profile_uri}",
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                tmpdir,
                str(source_path),
            ]
            subprocess.run(cmd, capture_output=True, timeout=60)
            tmp_docx = Path(tmpdir) / f"{source_path.stem}.docx"
            if tmp_docx.exists():
                return docx_to_markdown(tmp_docx, output_path)

    if source_path.suffix.lower() == ".odt" and zipfile.is_zipfile(source_path):
        with zipfile.ZipFile(source_path, "r") as zf:
            xml_bytes = zf.read("content.xml")
        root = ET.fromstring(xml_bytes)
        paras: list[str] = []
        for elem in root.iter():
            tag = elem.tag.split("}")[-1]
            if tag == "h":
                txt = "".join(elem.itertext()).strip()
                if txt:
                    paras.append(f"## {txt}")
            elif tag == "p":
                txt = "".join(elem.itertext()).strip()
                if txt:
                    paras.append(txt)
        output_path.write_text("\n\n".join(paras).strip() + "\n", encoding="utf-8")
        return output_path

    raw = source_path.read_text(encoding="utf-8", errors="replace")
    if source_path.suffix.lower() == ".rtf":
        raw = re.sub(r"\\[a-z]+\d* ?|[{}]", "", raw)
    output_path.write_text(raw.strip() + "\n", encoding="utf-8")
    return output_path
