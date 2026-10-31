import html
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any
from docx import Document
import pymupdf
from converter.markdown_engine import (
    docx_to_markdown,
    html_to_markdown,
    odt_or_office_to_markdown,
    pdf_to_markdown,
)


def _render_vector_pdf(*args: Any, **kwargs: Any) -> Path:
    """Renders text or line lists into a clean multi-page vector PDF (compatible with all caller signatures)."""
    output_path: Path | None = kwargs.get("output_path")
    title: str = str(kwargs.get("title", ""))
    content: Any = kwargs.get("lines") or kwargs.get("text")

    for arg in args:
        if isinstance(arg, Path) or (isinstance(arg, str) and arg.lower().endswith(".pdf") and "\n" not in arg):
            if output_path is None:
                output_path = Path(arg)
                continue
        if content is None:
            content = arg
        elif not title and isinstance(arg, str):
            title = arg

    if output_path is None:
        raise ValueError("output_path is required for _render_vector_pdf")

    if isinstance(content, list):
        lines = [str(line) for line in content]
    elif isinstance(content, str):
        lines = content.splitlines()
    else:
        lines = [str(content or "")]

    if title and (not lines or lines[0] != title):
        lines = [title, "-" * min(72, max(12, len(title))), ""] + lines

    pdf_doc = pymupdf.open()
    chunk_size = 46
    for i in range(0, max(1, len(lines)), chunk_size):
        page = pdf_doc.new_page()
        rect = pymupdf.Rect(40, 40, page.rect.width - 40, page.rect.height - 40)
        chunk_text = "\n".join(lines[i : i + chunk_size])
        page.insert_textbox(rect, chunk_text, fontsize=10, fontname="helv")

    pdf_doc.save(str(output_path))
    pdf_doc.close()
    return output_path


def _markdown_to_html(md_text: str) -> str:
    lines = md_text.splitlines()
    html_out: list[str] = []
    in_code = False
    in_ul = False
    in_table = False

    def _inline(s: str) -> str:
        s = html.escape(s)
        s = re.sub(r"`([^`]+)`", r"<code>\1</code>", s)
        s = re.sub(r"\*\*\*([^*]+)\*\*\*", r"<strong><em>\1</em></strong>", s)
        s = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", s)
        s = re.sub(r"\*([^*]+)\*", r"<em>\1</em>", s)
        s = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', s)
        return s

    for line in lines:
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                html_out.append("</code></pre>")
                in_code = False
            else:
                html_out.append("<pre><code>")
                in_code = True
            continue
        if in_code:
            html_out.append(html.escape(line))
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(set(c) <= {"-", ":", " "} for c in cells if c):
                continue
            if not in_table:
                html_out.append("<table><thead><tr>" + "".join(f"<th>{_inline(c)}</th>" for c in cells) + "</tr></thead><tbody>")
                in_table = True
            else:
                html_out.append("<tr>" + "".join(f"<td>{_inline(c)}</td>" for c in cells) + "</tr>")
            continue
        elif in_table:
            html_out.append("</tbody></table>")
            in_table = False

        if stripped.startswith("- ") or stripped.startswith("* "):
            if not in_ul:
                html_out.append("<ul>")
                in_ul = True
            html_out.append(f"<li>{_inline(stripped[2:])}</li>")
            continue
        elif in_ul:
            html_out.append("</ul>")
            in_ul = False

        if not stripped:
            continue
        if stripped.startswith("#"):
            m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
            if m:
                lvl = len(m.group(1))
                html_out.append(f"<h{lvl}>{_inline(m.group(2))}</h{lvl}>")
                continue
        if stripped.startswith("> "):
            html_out.append(f"<blockquote>{_inline(stripped[2:])}</blockquote>")
            continue
        html_out.append(f"<p>{_inline(stripped)}</p>")

    if in_ul:
        html_out.append("</ul>")
    if in_table:
        html_out.append("</tbody></table>")
    if in_code:
        html_out.append("</code></pre>")

    body = "\n".join(html_out)
    return (
        "<!DOCTYPE html><html><head><meta charset='utf-8'>"
        "<style>body{font-family:sans-serif;line-height:1.5;margin:32px;}"
        "table{border-collapse:collapse;width:100%;margin:12px 0;}"
        "th,td{border:1px solid #ccc;padding:6px 10px;text-align:left;}"
        "pre,code{background:#f4f4f5;padding:2px 5px;border-radius:4px;}</style>"
        f"</head><body>\n{body}\n</body></html>"
    )


def _markdown_to_docx(md_text: str, output_path: Path) -> Path:
    doc = Document()
    lines = md_text.splitlines()
    table_buffer: list[list[str]] = []

    def _flush_table() -> None:
        nonlocal table_buffer
        if not table_buffer:
            return
        cols = max(len(r) for r in table_buffer)
        tbl = doc.add_table(rows=len(table_buffer), cols=cols)
        for r_idx, row in enumerate(table_buffer):
            for c_idx, val in enumerate(row):
                tbl.cell(r_idx, c_idx).text = val
        table_buffer = []

    for line in lines:
        s = line.strip()
        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(set(c) <= {"-", ":", " "} for c in cells if c):
                continue
            table_buffer.append(cells)
            continue
        else:
            _flush_table()

        if not s or s == "---":
            continue
        m_h = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m_h:
            doc.add_heading(m_h.group(2), level=min(4, len(m_h.group(1))))
        elif s.startswith("- ") or s.startswith("* "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif re.match(r"^\d+[\.\)]\s+", s):
            doc.add_paragraph(re.sub(r"^\d+[\.\)]\s+", "", s), style="List Number")
        else:
            clean_s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
            doc.add_paragraph(clean_s)

    _flush_table()
    doc.save(str(output_path))
    return output_path


def convert_document(source_path: Path, output_path: Path, target_format: str) -> Path:
    src_ext = source_path.suffix.lower().lstrip(".")
    tgt_ext = target_format.lower().lstrip(".")

    if tgt_ext == "md":
        if src_ext == "pdf":
            return pdf_to_markdown(source_path, output_path)
        if src_ext == "docx":
            return docx_to_markdown(source_path, output_path)
        if src_ext in {"html", "htm"}:
            raw_html = source_path.read_text(encoding="utf-8", errors="replace")
            output_path.write_text(html_to_markdown(raw_html), encoding="utf-8")
            return output_path
        if src_ext in {"doc", "odt", "rtf"}:
            return odt_or_office_to_markdown(source_path, output_path)
        raw_txt = source_path.read_text(encoding="utf-8", errors="replace")
        output_path.write_text(raw_txt, encoding="utf-8")
        return output_path

    if tgt_ext == "txt":
        if src_ext == "docx":
            doc = Document(str(source_path))
            output_path.write_text("\n\n".join(p.text for p in doc.paragraphs if p.text.strip()), encoding="utf-8")
            return output_path
        if src_ext in {"html", "htm"}:
            raw_html = source_path.read_text(encoding="utf-8", errors="replace")
            output_path.write_text(html_to_markdown(raw_html), encoding="utf-8")
            return output_path
        output_path.write_text(source_path.read_text(encoding="utf-8", errors="replace"), encoding="utf-8")
        return output_path

    if tgt_ext == "html":
        if src_ext == "docx":
            with tempfile.NamedTemporaryFile(suffix=".md", delete=False) as tmp:
                tmp_md = Path(tmp.name)
            docx_to_markdown(source_path, tmp_md)
            md_str = tmp_md.read_text(encoding="utf-8")
            tmp_md.unlink(missing_ok=True)
            output_path.write_text(_markdown_to_html(md_str), encoding="utf-8")
            return output_path
        raw_text = source_path.read_text(encoding="utf-8", errors="replace")
        output_path.write_text(_markdown_to_html(raw_text), encoding="utf-8")
        return output_path

    if tgt_ext == "docx":
        if src_ext in {"doc", "odt", "rtf"}:
            soffice = shutil.which("libreoffice") or shutil.which("soffice")
            if soffice:
                with tempfile.TemporaryDirectory() as tmpdir:
                    profile_uri = Path(tmpdir).as_uri() + "/lo_profile"
                    subprocess.run(
                        [soffice, f"-env:UserInstallation={profile_uri}", "--headless", "--convert-to", "docx", "--outdir", tmpdir, str(source_path)],
                        capture_output=True,
                        timeout=60,
                    )
                    gen = Path(tmpdir) / f"{source_path.stem}.docx"
                    if gen.exists():
                        shutil.move(str(gen), str(output_path))
                        return output_path
        if src_ext in {"html", "htm"}:
            md_str = html_to_markdown(source_path.read_text(encoding="utf-8", errors="replace"))
            return _markdown_to_docx(md_str, output_path)
        md_str = source_path.read_text(encoding="utf-8", errors="replace")
        return _markdown_to_docx(md_str, output_path)

    if tgt_ext == "pdf":
        soffice = shutil.which("libreoffice") or shutil.which("soffice")
        if src_ext in {"docx", "doc", "odt", "rtf"} and soffice:
            with tempfile.TemporaryDirectory() as tmpdir:
                profile_uri = Path(tmpdir).as_uri() + "/lo_profile"
                subprocess.run(
                    [soffice, f"-env:UserInstallation={profile_uri}", "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(source_path)],
                    capture_output=True,
                    timeout=60,
                )
                gen = Path(tmpdir) / f"{source_path.stem}.pdf"
                if gen.exists():
                    shutil.move(str(gen), str(output_path))
                    return output_path

        if src_ext == "docx":
            with tempfile.NamedTemporaryFile(suffix=".md", delete=False) as tmp:
                tmp_md = Path(tmp.name)
            docx_to_markdown(source_path, tmp_md)
            text_body = tmp_md.read_text(encoding="utf-8")
            tmp_md.unlink(missing_ok=True)
        else:
            text_body = source_path.read_text(encoding="utf-8", errors="replace")

        return _render_vector_pdf(text_body.splitlines(), output_path)

    shutil.copy2(source_path, output_path)
    return output_path
