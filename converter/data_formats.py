import csv
import json
import shutil
import subprocess
import tempfile
import tomllib
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any
import openpyxl
from openpyxl.styles import Font
import tomli_w
import yaml
from converter.documents import _render_vector_pdf


def _load_structured_data(source_path: Path) -> Any:
    ext = source_path.suffix.lower().lstrip(".")

    if ext == "json":
        return json.loads(source_path.read_text(encoding="utf-8"))

    if ext in {"yaml", "yml"}:
        return yaml.safe_load(source_path.read_text(encoding="utf-8"))

    if ext == "toml":
        return tomllib.loads(source_path.read_text(encoding="utf-8"))

    if ext in {"csv", "tsv"}:
        delim = "\t" if ext == "tsv" else ","
        with source_path.open("r", encoding="utf-8", errors="replace", newline="") as f:
            reader = csv.DictReader(f, delimiter=delim)
            return list(reader)

    if ext in {"xlsx", "xls", "ods"}:
        actual_xlsx = source_path
        if ext in {"xls", "ods"}:
            soffice = shutil.which("libreoffice") or shutil.which("soffice")
            if not soffice:
                raise RuntimeError(f"LibreOffice is required to read .{ext} spreadsheets.")
            tmpdir = tempfile.mkdtemp()
            cmd = [soffice, "--headless", "--convert-to", "xlsx", "--outdir", tmpdir, str(source_path)]
            subprocess.run(cmd, capture_output=True, timeout=60)
            actual_xlsx = Path(tmpdir) / f"{source_path.stem}.xlsx"

        wb = openpyxl.load_workbook(actual_xlsx, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        records = []
        for row in rows[1:]:
            if any(cell is not None for cell in row):
                records.append({headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))})
        return records

    if ext == "xml":
        tree = ET.parse(source_path)
        root = tree.getroot()
        records = []
        for child in root:
            item = dict(child.attrib)
            for sub in child:
                item[sub.tag] = sub.text
            if not item and child.text and child.text.strip():
                item["value"] = child.text.strip()
            records.append(item)
        return records

    raise ValueError(f"Unsupported data input format: .{ext}")


def _normalize_to_rows(data: Any) -> list[dict[str, Any]]:
    if isinstance(data, list):
        if all(isinstance(x, dict) for x in data):
            return data
        return [{"value": x} for x in data]
    if isinstance(data, dict):
        for val in data.values():
            if isinstance(val, list) and val and all(isinstance(x, dict) for x in val):
                return val
        return [data]
    return [{"value": data}]


def convert_data(source_path: Path, output_path: Path, target_format: str) -> Path:
    tgt = target_format.lower().lstrip(".")
    src_ext = source_path.suffix.lower().lstrip(".")

    # Direct LibreOffice Calc export for ODS / PDF when LibreOffice is installed
    if tgt in {"ods", "pdf"} and src_ext in {"xlsx", "xls", "ods", "csv", "tsv"}:
        soffice = shutil.which("libreoffice") or shutil.which("soffice")
        if soffice:
            with tempfile.TemporaryDirectory() as tmpdir:
                profile_uri = Path(tmpdir).as_uri() + "/lo_profile"
                cmd = [
                    soffice, f"-env:UserInstallation={profile_uri}",
                    "--headless", "--convert-to", tgt, "--outdir", tmpdir, str(source_path)
                ]
                subprocess.run(cmd, capture_output=True, timeout=60)
                gen = Path(tmpdir) / f"{source_path.stem}.{tgt}"
                if gen.exists():
                    shutil.move(str(gen), str(output_path))
                    return output_path

    data = _load_structured_data(source_path)

    if tgt == "json":
        output_path.write_text(json.dumps(data, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
        return output_path

    if tgt in {"yaml", "yml"}:
        output_path.write_text(yaml.safe_dump(data, sort_keys=False, allow_unicode=True), encoding="utf-8")
        return output_path

    if tgt == "toml":
        toml_payload = data if isinstance(data, dict) else {"items": data}
        clean_payload = json.loads(json.dumps(toml_payload, default=str))
        output_path.write_text(tomli_w.dumps(clean_payload), encoding="utf-8")
        return output_path

    if tgt in {"csv", "tsv"}:
        delim = "\t" if tgt == "tsv" else ","
        rows = _normalize_to_rows(data)
        headers: list[str] = []
        for r in rows:
            for k in r.keys():
                if str(k) not in headers:
                    headers.append(str(k))
        with output_path.open("w", encoding="utf-8", newline="") as f:
            if headers:
                writer = csv.DictWriter(f, fieldnames=headers, delimiter=delim)
                writer.writeheader()
                for r in rows:
                    writer.writerow({str(k): v for k, v in r.items()})
        return output_path

    if tgt == "xlsx":
        rows = _normalize_to_rows(data)
        headers = []
        for r in rows:
            for k in r.keys():
                if str(k) not in headers:
                    headers.append(str(k))
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        if headers:
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for r in rows:
                ws.append([r.get(h, "") for h in headers])
        wb.save(output_path)
        return output_path

    if tgt == "pdf":
        rows = _normalize_to_rows(data)
        lines = [" | ".join(f"{k}: {v}" for k, v in r.items()) for r in rows]
        return _render_vector_pdf("\n".join(lines), output_path)

    if tgt == "xml":
        rows = _normalize_to_rows(data)
        root = ET.Element("dataset")
        for r in rows:
            item_el = ET.SubElement(root, "record")
            for k, v in r.items():
                tag_name = "".join(c if c.isalnum() or c == "_" else "_" for c in str(k))
                if not tag_name or tag_name[0].isdigit():
                    tag_name = f"field_{tag_name}"
                sub = ET.SubElement(item_el, tag_name)
                sub.text = "" if v is None else str(v)
        ET.indent(root, space="  ")
        ET.ElementTree(root).write(output_path, encoding="utf-8", xml_declaration=True)
        return output_path

    raise ValueError(f"Unsupported data target format: {target_format}")
